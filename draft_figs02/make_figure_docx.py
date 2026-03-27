import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from docx.shared import Pt

images_list = [
    "fig01_p.compare.std_mean.l2pVSxtrack_map.png",
    "fig02_p.omni_uvVStrack_guv_c.ek.nearest.selected.png",
    "fig03_p.scatter_omniVSrama_uVSgu_ww_ek.png",
    "fig04_p.scatter_drifter_ve_vs_gu_ww_ek.png",
    "fig05_p.radar_uvVStrack_guv.at.intersections_single.png",
    "fig06_p.gc.clim.cross_structure_x_arrows_map_ERS.png",
    "fig07_p.gc.clim.lat_wise.along_xwidth.png",
    "fig08_p.along_shore_at_155_116.intersection.S6A.png",
    "fig09_p.speed_at_coastal_points.eicc.trends.all_in_one.ERS.png",
]

for image in images_list:
    if not os.path.exists(image):
        raise FileNotFoundError(f"Image {image} not found.")
print("all good")

captions_list = [
    "captions/p.compare.std_mean.l2pVSxtrack_map.txt",
    "captions/p.omni_uvVStrack_guv_c.ek.nearest.selected.txt",
    "captions/p.scatter_omniVSrama_uVSgu_ww_ek.txt",
    "captions/p.scatter_drifter_ve_vs_gu_ww_ek.txt",
    "captions/p.radar_uvVStrack_guv.at.intersections_single.txt",
    "captions/p.gc.clim.cross_structure_x_arrows_map_ERS.txt",
    "captions/p.gc.clim.lat_wise.along_xwidth.txt",
    "captions/p.along_shore_at_155_116.intersection.S6A.txt",
    "captions/p.speed_at_coastal_points.eicc.trends.all_in_one.ERS.txt",
]

# Caption styling options (edit these as needed).
CAPTION_FONT_NAME = "Times New Roman"
CAPTION_FONT_SIZE_PT = 12
CAPTION_ALIGNMENT = "center"  # one of: left, center, right, justify

ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

doc_name = "figures_v2.1.docx"
if os.path.exists(doc_name):
    os.remove(doc_name)

doc = Document()
inserted_count = 0

for i, img in enumerate(images_list):
    if not os.path.exists(img):
        print(f"Skipping missing image: {img}")
        continue

    doc.add_picture(img, width=Inches(6))
    inserted_count += 1

    caption_text = ""
    if i < len(captions_list):
        caption_path = captions_list[i]
        if os.path.exists(caption_path):
            with open(caption_path, "r", encoding="utf-8") as f:
                caption_text = f.read().strip()

    prefix = f"Fig{inserted_count}"
    full_caption = f"{prefix}: {caption_text}" if caption_text else prefix

    paragraph = doc.add_paragraph()
    paragraph.alignment = ALIGNMENT_MAP.get(
        CAPTION_ALIGNMENT.lower(), WD_ALIGN_PARAGRAPH.CENTER
    )
    run = paragraph.add_run(full_caption)
    run.font.name = CAPTION_FONT_NAME
    run.font.size = Pt(CAPTION_FONT_SIZE_PT)

    if i < len(images_list) - 1:
        doc.add_page_break()

doc.save(doc_name)
print(f"Inserted {inserted_count} figure(s) into {doc_name}")
