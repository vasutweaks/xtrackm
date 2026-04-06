import os
import re
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
# codex resume 019d644d-ce0b-7e22-af34-8fb378512249
image_pattern = re.compile(r'^(fig|sup)(\d+)_.*\.png$')

# Caption styling options (edit these as needed).
CAPTION_FONT_NAME = "Times New Roman"
CAPTION_FONT_SIZE_PT = 12
CAPTION_ALIGNMENT = "justify"  # one of: left, center, right, justify

ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

doc_name = "figures_v2.1.docx"
if os.path.exists(doc_name):
    os.remove(doc_name)

image_entries = []
for filename in os.listdir("."):
    m = image_pattern.match(filename)
    if not m:
        continue
    group_kind, group_num = m.group(1), int(m.group(2))
    image_entries.append((group_kind, group_num, filename))

# Deterministic order: all figures first, then supplements; each by numeric index.
image_entries.sort(key=lambda x: (0 if x[0] == "fig" else 1, x[1], x[2]))

doc = Document()
inserted_count = 0
section = doc.sections[0]
usable_width = section.page_width - section.left_margin - section.right_margin

for i, (group_kind, group_num, img) in enumerate(image_entries):
    if not os.path.exists(img):
        print(f"Skipping missing image: {img}")
        continue

    doc.add_picture(img, width=usable_width)
    inserted_count += 1

    caption_text = ""
    caption_path = f"{os.path.splitext(img)[0]}.txt"
    if os.path.exists(caption_path):
        with open(caption_path, "r", encoding="utf-8") as f:
            caption_text = f.read().strip()

    prefix = f"Fig{group_num}" if group_kind == "fig" else f"Sup Fig{group_num}"
    full_caption = f"{prefix}: {caption_text}" if caption_text else prefix

    paragraph = doc.add_paragraph()
    paragraph.alignment = ALIGNMENT_MAP.get(
        CAPTION_ALIGNMENT.lower(), WD_ALIGN_PARAGRAPH.CENTER
    )
    run = paragraph.add_run(full_caption)
    run.font.name = CAPTION_FONT_NAME
    run.font.size = Pt(CAPTION_FONT_SIZE_PT)

    if i < len(image_entries) - 1:
        doc.add_page_break()

doc.save(doc_name)
print(f"Inserted {inserted_count} figure(s) into {doc_name}")
# convert docx to pdf
try:
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", doc_name],
        check=True,
    )
except (OSError, subprocess.CalledProcessError) as exc:
    print(f"PDF conversion failed: {exc}")
